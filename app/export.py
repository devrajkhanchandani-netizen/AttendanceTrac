from datetime import date
from pathlib import Path
from typing import Optional

from openpyxl import Workbook
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from app.attendance import ABSENT, NEEDS_REVIEW, PRESENT, AttendanceReport

MARKS = {PRESENT: "P", ABSENT: "A", NEEDS_REVIEW: "?"}

HEADER_FILL = PatternFill("solid", start_color="D9D9D9")
BOLD = Font(bold=True)


def _text_cell(sheet, row: int, column: int, value):
    text = "" if value is None else str(value)
    cell = sheet.cell(row=row, column=column, value=text)
    cell.data_type = "s"
    return cell


def _write_header(sheet, row: int, titles: list) -> None:
    for col, title in enumerate(titles, start=1):
        cell = sheet.cell(row=row, column=col, value=title)
        cell.font = BOLD
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal="center")


def _set_widths(sheet, widths: list) -> None:
    for col, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(col)].width = width


def export_attendance(report: AttendanceReport,
                      path,
                      class_name: str = "",
                      session_date: Optional[date] = None,
                      photo_name: str = "") -> Path:
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    session_date = session_date or date.today()

    workbook = Workbook()

    sheet = workbook.active
    sheet.title = "Attendance"

    sheet["A1"] = "AttendanceTrac - Attendance Report"
    sheet["A1"].font = Font(bold=True, size=14)
    sheet["A2"] = "Class:"
    _text_cell(sheet, 2, 2, class_name)
    sheet["A3"] = "Date:"
    sheet["B3"] = session_date.isoformat()
    sheet["A4"] = "Photo:"
    _text_cell(sheet, 4, 2, photo_name)
    for cell in ("A2", "A3", "A4"):
        sheet[cell].font = BOLD

    header_row = 6
    _write_header(sheet, header_row,
                  ["No.", "Student", "Mark", "Status", "Confidence"])

    row = header_row
    for number, record in enumerate(report.records, start=1):
        row += 1
        sheet.cell(row=row, column=1, value=number)
        _text_cell(sheet, row, 2, record.student)
        mark = sheet.cell(row=row, column=3, value=MARKS.get(record.status, "?"))
        mark.alignment = Alignment(horizontal="center")
        mark.font = BOLD
        sheet.cell(row=row, column=4, value=record.status)
        confidence = sheet.cell(row=row, column=5)
        if record.confidence is not None:
            confidence.value = round(record.confidence, 2)
            confidence.number_format = "0.00"

    first, last = header_row + 1, max(row, header_row + 1)

    mark_range = f"C{first}:C{last}"
    for letter, colour in (("P", "C6EFCE"), ("A", "FFC7CE"), ("?", "FFEB9C")):
        sheet.conditional_formatting.add(
            mark_range,
            CellIsRule(operator="equal", formula=[f'"{letter}"'],
                       fill=PatternFill("solid", start_color=colour,
                                        end_color=colour)),
        )

    counts = {status: 0 for status in MARKS}
    for record in report.records:
        counts[record.status] = counts.get(record.status, 0) + 1
    summary_row = row + 2
    sheet.cell(row=summary_row, column=2, value="Total students").font = BOLD
    sheet.cell(row=summary_row, column=3, value=len(report.records))
    for offset, status in enumerate((PRESENT, ABSENT, NEEDS_REVIEW), start=1):
        sheet.cell(row=summary_row + offset, column=2, value=status).font = BOLD
        sheet.cell(row=summary_row + offset, column=3, value=counts[status])

    _set_widths(sheet, [6, 22, 8, 16, 12])
    sheet.freeze_panes = sheet.cell(row=header_row + 1, column=1)

    review = workbook.create_sheet("Needs Review")
    _write_header(review, 1,
                  ["Face #", "Possible student", "Similarity", "Reason"])
    if report.flagged_faces:
        for r, face in enumerate(report.flagged_faces, start=2):
            review.cell(row=r, column=1, value=face.face_index)
            _text_cell(review, r, 2, face.best_identity)
            sim = review.cell(row=r, column=3, value=round(face.similarity, 2))
            sim.number_format = "0.00"
            review.cell(row=r, column=4, value=face.reason)
    else:
        review.cell(row=2, column=1, value="No faces need review.")
    _set_widths(review, [8, 20, 12, 40])
    review.freeze_panes = "A2"

    workbook.save(path)
    return path